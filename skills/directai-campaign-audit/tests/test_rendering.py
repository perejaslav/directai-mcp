#!/usr/bin/env python3
"""Отчёты: единый источник, форматы, отказ при FAIL (ТЗ §11, §13)."""
from __future__ import annotations

import datetime
import io
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))

import factory  # noqa: E402
from auditlib import build, render_html, render_md, render_xlsx  # noqa: E402
from auditlib.model import fmt_money  # noqa: E402

TODAY = datetime.date(2026, 10, 23)


class RenderingCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="audit-render-"))
        self.root = factory.make_session(
            self.tmp / "s",
            modes=["deep"],
            queries=[("окна даром", dict(impressions=2000, clicks=60, cost=2400, conversions=0)),
                     ("купить окна пвх", dict(impressions=5000, clicks=150, cost=9000, conversions=9))],
            placements=[("bad.example", dict(impressions=80000, clicks=500, cost=3200, conversions=0))],
            devices=[("MOBILE", dict(impressions=50000, clicks=1500, cost=95000, conversions=10)),
                     ("DESKTOP", dict(impressions=50000, clicks=500, cost=25000, conversions=30))],
            metrika=dict(visits=1500),
            ads=[("1923568665500247917", ["DisplayUrlPath"])],
            moderation=[("1923568665500247917", "REJECTED", "превышение лимита заглавных")],
        )
        self.audit = build.build(self.root, today=TODAY)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)


class TestMarkdown(RenderingCase):
    def test_has_all_required_sections(self):
        text = render_md.render(self.audit)
        for index in range(1, 11):
            self.assertIn(f"## {index}.", text, f"нет раздела {index}")

    def test_summary_has_at_most_five_points(self):
        text = render_md.render(self.audit)
        block = text.split("## 2. Краткое резюме")[1].split("## 3.")[0]
        points = [line for line in block.splitlines() if line.startswith("- ")]
        self.assertLessEqual(len(points), 5)

    def test_numbers_match_json(self):
        text = render_md.render(self.audit)
        metrics = self.audit["metrics"]
        self.assertIn(fmt_money(metrics["a"]["cost"]), text)
        self.assertIn(fmt_money(metrics["b"]["cost"]), text)
        self.assertIn(str(metrics["a"]["clicks"]).replace(",", " ") if metrics["a"]["clicks"] >= 1000
                      else str(metrics["a"]["clicks"]), text)

    def test_formula_is_explained(self):
        text = render_md.render(self.audit)
        self.assertIn("Δ = A − B", text)
        self.assertIn("Δ% = (A − B) / B × 100", text)

    def test_no_invented_values(self):
        text = render_md.render(self.audit)
        self.assertNotIn("None", text)
        self.assertNotIn("nan", text.lower())

    def test_each_recommendation_has_verification(self):
        text = render_md.render(self.audit)
        for rec in self.audit["recommendations"]:
            self.assertIn(rec["id"], text)
            self.assertIn("Как проверить", text)


class TestExcel(RenderingCase):
    def load(self) -> dict:
        from openpyxl import load_workbook
        data = render_xlsx.render(self.audit)
        self.assertGreater(len(data), 5000)
        return load_workbook(io.BytesIO(data))

    def test_all_sheets_present(self):
        workbook = self.load()
        expected = ["Обзор", "Динамика", "Кампании и группы", "Ключевые фразы",
                    "Поисковые запросы", "Площадки РСЯ", "Устройства и регионы",
                    "Рекомендации", "Источники и ограничения"]
        self.assertEqual(workbook.sheetnames, expected)

    def test_header_frozen_and_filtered(self):
        workbook = self.load()
        dynamics = workbook["Динамика"]
        self.assertEqual(dynamics.freeze_panes, "A2")
        self.assertTrue(dynamics.auto_filter.ref)
        recommendations = workbook["Рекомендации"]
        self.assertEqual(recommendations.freeze_panes, "A2")

    def test_numbers_stay_numbers(self):
        workbook = self.load()
        dynamics = workbook["Динамика"]
        self.assertIsInstance(dynamics["B2"].value, int)
        self.assertIsInstance(dynamics["B5"].value, (int, float))
        self.assertIn("₽", dynamics["B5"].number_format)

    def test_ids_are_text(self):
        workbook = self.load()
        overview = workbook["Обзор"]
        campaign_cell = next(row for row in overview.iter_rows(min_row=2) if row[0].value == "Кампания ID")[1]
        self.assertEqual(campaign_cell.number_format, "@")
        self.assertIsInstance(campaign_cell.value, str)

    def test_priority_is_highlighted(self):
        workbook = self.load()
        sheet = workbook["Рекомендации"]
        if sheet.max_row > 1:
            self.assertIsNotNone(sheet["B2"].fill.fgColor.rgb)

    def test_empty_sheet_is_marked(self):
        workbook = self.load()
        keywords = workbook["Ключевые фразы"]
        self.assertEqual(keywords["A1"].value, "Данные отсутствуют")


class TestHtml(RenderingCase):
    def test_self_contained(self):
        page = render_html.render(self.audit)
        self.assertIn("<!DOCTYPE html>", page)
        self.assertNotIn("http://", page)
        self.assertNotIn("https://", page)
        self.assertNotIn("<link", page)
        self.assertNotIn("src=\"http", page)

    def test_contains_kpi_and_recommendations(self):
        page = render_html.render(self.audit)
        self.assertIn(fmt_money(self.audit["metrics"]["a"]["cost"]), page)
        for rec in self.audit["recommendations"][:3]:
            self.assertIn(rec["id"] if rec["id"] in page else rec["problem"], page)

    def test_has_filters_and_tabs(self):
        page = render_html.render(self.audit)
        self.assertIn('data-filter="P0"', page)
        self.assertIn('data-tab="devices"', page)
        self.assertIn("addEventListener", page)

    def test_missing_data_block(self):
        page = render_html.render(self.audit)
        self.assertIn("Ограничения и пробелы в данных", page)
        self.assertIn("Источники", page)

    def test_hypotheses_highlighted(self):
        page = render_html.render(self.audit)
        self.assertIn("Непроверенные гипотезы", page)


class TestSingleSource(RenderingCase):
    def test_all_formats_use_same_numbers(self):
        metrics = self.audit["metrics"]["a"]
        md = render_md.render(self.audit)
        page = render_html.render(self.audit)
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(render_xlsx.render(self.audit)))
        cell = workbook["Обзор"]
        excel_cost = next(row[1].value for row in cell.iter_rows(min_row=2) if row[0].value == "Расход A, ₽")
        self.assertAlmostEqual(excel_cost, metrics["cost"])
        self.assertIn(fmt_money(metrics["cost"]), md)
        self.assertIn(fmt_money(metrics["cost"]), page)

    def test_change_in_json_changes_all_views(self):
        mutated = json.loads(json.dumps(self.audit))
        mutated["metrics"]["a"]["cost"] = 777777.77
        mutated["metrics"]["comparison"]["cost"] = {"a": 777777.77, "b": 118000.0,
                                                    "kind": "money", "available": True,
                                                    "abs": 659777.77, "pct": 559.13, "pp": None,
                                                    "base_zero": False, "display_hides_change": False,
                                                    "note": None}
        md = render_md.render(mutated)
        page = render_html.render(mutated)
        self.assertIn("777 777,77 ₽", md)
        self.assertIn("777 777,77 ₽", page)


class TestPublishGuard(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="audit-guard-"))
        root = factory.make_session(self.tmp / "s", modes=["deep"])
        audit = build.build(root, today=TODAY)
        audit["validation"] = {"status": "FAIL", "issues": [], "summary": {"errors": 1, "warnings": 0}}
        self.path = self.tmp / "audit.json"
        self.path.write_text(json.dumps(audit, ensure_ascii=False), encoding="utf-8")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_markdown_refuses_fail(self):
        result = subprocess.run([sys.executable, str(ROOT / "scripts" / "render_md.py"), str(self.path)],
                                capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 2)
        self.assertIn("ОТКАЗ", result.stdout)
        self.assertFalse(self.path.with_suffix(".md").exists())

    def test_diagnostic_flag_allows_report(self):
        result = subprocess.run([sys.executable, str(ROOT / "scripts" / "render_md.py"), str(self.path),
                                 "--diagnostic"], capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0)
        text = self.path.with_suffix(".md").read_text(encoding="utf-8")
        self.assertIn("Диагностический отчёт", text)

    def test_validate_cli_exit_codes(self):
        broken = json.loads(self.path.read_text(encoding="utf-8"))
        broken["metrics"]["comparison"]["cost"]["abs"] = 999999.0  # Δ ≠ A−B
        path = self.tmp / "broken.json"
        path.write_text(json.dumps(broken, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run([sys.executable, str(ROOT / "scripts" / "validate_audit.py"), str(path)],
                                capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 2)
        self.assertIn("FAIL", result.stdout)
        self.assertIn("CMP-02", result.stdout)


if __name__ == "__main__":
    unittest.main()
