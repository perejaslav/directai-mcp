#!/usr/bin/env python3
"""Excel-отчёт из audit.json (ТЗ §11.2).

Требования: фиксированная шапка, автофильтры, числовые форматы, ID текстом,
выделение приоритетов, корректное открытие в Microsoft Excel.
"""
from __future__ import annotations

from .model import fmt_id

MONEY = '#,##0.00 "₽"'
PERCENT = '#,##0.00"%"'
NUMBER = '#,##0.00'
INTEGER = '#,##0'
TEXT = '@'

PRIORITY_FILL = {
    "P0": "F8D7DA",
    "P1": "FCE5CD",
    "P2": "FFF2CC",
    "P3": "EDEDED",
}

SHEETS = ("Обзор", "Динамика", "Кампании и группы", "Ключевые фразы", "Поисковые запросы",
          "Площадки РСЯ", "Устройства и регионы", "Рекомендации", "Источники и ограничения")


def _sheet(wb, title):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(title)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    return ws


def _header(ws, columns, row=1):
    from openpyxl.styles import Alignment, Font, PatternFill
    for index, name in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=index, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4A5568")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def _widths(ws, widths):
    from openpyxl.utils import get_column_letter
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_table(ws, columns, rows, *, formats=None, widths=None, autofilter=True):
    from openpyxl.styles import Alignment
    _header(ws, columns)
    formats = formats or {}
    for r_index, row in enumerate(rows, start=2):
        for c_index, name in enumerate(columns, start=1):
            value = row.get(name)
            cell = ws.cell(row=r_index, column=c_index, value=value)
            fmt = formats.get(name)
            if fmt:
                cell.number_format = fmt
            if name in ("ID", "Id", "Кампания", "Объект ID", "Finding") or fmt == TEXT:
                cell.number_format = TEXT
                cell.value = str(value) if value is not None else None
            cell.alignment = Alignment(vertical="top", wrap_text=name in
                                       ("Проблема", "Что сделать", "Доказательства", "Риск", "Проверка"))
    if rows and autofilter:
        from openpyxl.utils import get_column_letter
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    if widths:
        _widths(ws, widths)


def _empty(ws, reason):
    from openpyxl.styles import Font
    ws.cell(row=1, column=1, value="Данные отсутствуют").font = Font(bold=True)
    ws.cell(row=2, column=1, value=reason or "источник не получен или режим не применим")


def render(audit: dict) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("нужен openpyxl: pip install openpyxl") from exc
    import io

    wb = Workbook()
    wb.remove(wb.active)
    scope = audit.get("scope") or {}
    periods = audit.get("periods") or {}
    metrics = audit.get("metrics") or {}
    dq = audit.get("data_quality") or {}

    # 1. Обзор
    ws = _sheet(wb, "Обзор")
    _header(ws, ["Показатель", "Значение"])
    overview = [
        ("audit_id", audit.get("audit_id")),
        ("Создан", audit.get("created_at")),
        ("Кабинет", scope.get("account_login") or scope.get("account")),
        ("Кампания ID", str(scope.get("campaign_id"))),
        ("Кампания", scope.get("campaign_name")),
        ("Тип", scope.get("campaign_type")),
        ("Состояние", scope.get("campaign_state")),
        ("Режимы", ", ".join(audit.get("mode") or [])),
        ("Период A", f"{periods.get('a', {}).get('from')} — {periods.get('a', {}).get('to')}"),
        ("Период B", f"{periods.get('b', {}).get('from')} — {periods.get('b', {}).get('to')}"),
        ("Основная цель", (audit.get("goal") or {}).get("label")),
        ("Атрибуция", (audit.get("attribution") or {}).get("effective")),
        ("Качество данных", dq.get("status")),
        ("Валидация", (audit.get("validation") or {}).get("status")),
        ("Расход A, ₽", metrics.get("a", {}).get("cost")),
        ("Конверсии A", metrics.get("a", {}).get("conversions")),
        ("CPA A, ₽", metrics.get("a", {}).get("cpa")),
        ("Расход под риском, ₽", (audit.get("spend_accounting") or {}).get("at_risk_spend", {}).get("sum")),
        ("Потенциальная экономия, ₽", (audit.get("spend_accounting") or {}).get("potential_saving", {}).get("value")),
    ]
    for r_index, (name, value) in enumerate(overview, start=2):
        ws.cell(row=r_index, column=1, value=name)
        cell = ws.cell(row=r_index, column=2, value=value)
        if isinstance(value, (int, float)):
            cell.number_format = MONEY if "₽" in name else NUMBER
        elif value is not None:
            cell.value = str(value)
            cell.number_format = TEXT
    _widths(ws, [32, 60])
    ws.cell(row=len(overview) + 3, column=1,
            value="Потенциальная экономия — оценка при условиях из листа «Рекомендации», не гарантия.").font = None

    # 2. Динамика
    ws = _sheet(wb, "Динамика")
    comparison = metrics.get("comparison") or {}
    rows = []
    labels = {"impressions": ("Показы", INTEGER), "clicks": ("Клики", INTEGER), "ctr": ("CTR, %", PERCENT),
              "cost": ("Расход, ₽", MONEY), "cpc": ("CPC, ₽", MONEY), "conversions": ("Конверсии", INTEGER),
              "cr": ("CR, %", PERCENT), "cpa": ("CPA, ₽", MONEY), "bounce_rate": ("Отказы, %", PERCENT)}
    for key, (label, fmt) in labels.items():
        item = comparison.get(key) or {}
        rows.append({
            "Метрика": label,
            "A": metrics.get("a", {}).get(key),
            "B": metrics.get("b", {}).get(key),
            "Δ": item.get("abs"),
            "Δ%": item.get("pct"),
            "п.п.": item.get("pp"),
            "База B = 0": "да" if item.get("base_zero") else "",
            "Примечание": item.get("note") or "",
        })
    formats = {"A": NUMBER, "B": NUMBER, "Δ": NUMBER, "Δ%": PERCENT, "п.п.": NUMBER}
    _write_table(ws, ["Метрика", "A", "B", "Δ", "Δ%", "п.п.", "База B = 0", "Примечание"], rows,
                 formats=formats, widths=[18, 14, 14, 14, 12, 12, 12, 50])
    for r_index, key in enumerate(labels, start=2):
        fmt = labels[key][1]
        for column in (2, 3, 4):
            ws.cell(row=r_index, column=column).number_format = fmt
        if labels[key][0].endswith("%"):
            ws.cell(row=r_index, column=5).number_format = PERCENT
    ws.cell(row=len(rows) + 3, column=1,
            value="Δ = A − B; Δ% = (A − B) / B × 100; п.п. — процентные пункты для долей.")

    # 3. Кампании и группы
    ws = _sheet(wb, "Кампании и группы")
    adgroups = audit.get("segment_rows", {}).get("adgroups") or []
    block = audit.get("segments", {}).get("adgroups") or {}
    if not adgroups:
        _empty(ws, block.get("note") or "stats_adgroups не получен")
    else:
        rows = [{
            "Группа": r.get("key"), "Показы": r.get("impressions"), "Клики": r.get("clicks"),
            "CTR, %": r.get("ctr"), "Расход, ₽": r.get("cost"), "Конверсии": r.get("conversions"),
            "CPA, ₽": r.get("cpa"), "CR, %": r.get("cr"),
        } for r in adgroups]
        _write_table(ws, ["Группа", "Показы", "Клики", "CTR, %", "Расход, ₽", "Конверсии", "CPA, ₽", "CR, %"],
                     rows, formats={"Показы": INTEGER, "Клики": INTEGER, "CTR, %": PERCENT,
                                    "Расход, ₽": MONEY, "Конверсии": INTEGER, "CPA, ₽": MONEY,
                                    "CR, %": PERCENT},
                     widths=[40, 12, 10, 10, 14, 12, 12, 10])

    # 4. Ключевые фразы
    ws = _sheet(wb, "Ключевые фразы")
    keywords = audit.get("segment_rows", {}).get("keywords") or []
    block = audit.get("segments", {}).get("keywords") or {}
    if not keywords:
        _empty(ws, block.get("note") or "stats_keywords не получен")
    else:
        rows = [{"Фраза": r.get("key"), "Показы": r.get("impressions"), "Клики": r.get("clicks"),
                 "CTR, %": r.get("ctr"), "Расход, ₽": r.get("cost"), "Конверсии": r.get("conversions"),
                 "CPA, ₽": r.get("cpa")} for r in keywords]
        _write_table(ws, ["Фраза", "Показы", "Клики", "CTR, %", "Расход, ₽", "Конверсии", "CPA, ₽"], rows,
                     formats={"Показы": INTEGER, "Клики": INTEGER, "CTR, %": PERCENT,
                              "Расход, ₽": MONEY, "Конверсии": INTEGER, "CPA, ₽": MONEY},
                     widths=[50, 12, 10, 10, 14, 12, 12])

    # 5. Поисковые запросы
    ws = _sheet(wb, "Поисковые запросы")
    queries = audit.get("segment_rows", {}).get("search_queries") or []
    block = audit.get("segments", {}).get("search_queries") or {}
    if not queries:
        _empty(ws, block.get("note") or "stats_search_queries не получен или строк нет")
    else:
        rows = [{"Запрос": r.get("key"), "Класс": r.get("intent") or "",
                 "Показы": r.get("impressions"), "Клики": r.get("clicks"), "Расход, ₽": r.get("cost"),
                 "Конверсии": r.get("conversions"), "CPA, ₽": r.get("cpa"),
                 "Основание класса": r.get("intent_reason") or ""} for r in queries]
        _write_table(ws, ["Запрос", "Класс", "Показы", "Клики", "Расход, ₽", "Конверсии", "CPA, ₽",
                          "Основание класса"], rows,
                     formats={"Показы": INTEGER, "Клики": INTEGER, "Расход, ₽": MONEY,
                              "Конверсии": INTEGER, "CPA, ₽": MONEY},
                     widths=[45, 18, 10, 10, 14, 12, 12, 45])

    # 6. Площадки РСЯ
    ws = _sheet(wb, "Площадки РСЯ")
    placements = audit.get("segment_rows", {}).get("placements") or []
    block = audit.get("segments", {}).get("placements") or {}
    if not placements:
        _empty(ws, block.get("note") or "строк нет — пустой ответ не доказывает отсутствие показов")
    else:
        rows = [{"Площадка": r.get("key"), "Показы": r.get("impressions"), "Клики": r.get("clicks"),
                 "CTR, %": r.get("ctr"), "Расход, ₽": r.get("cost"), "Конверсии": r.get("conversions"),
                 "CPA, ₽": r.get("cpa"), "CR, %": r.get("cr")} for r in placements]
        _write_table(ws, ["Площадка", "Показы", "Клики", "CTR, %", "Расход, ₽", "Конверсии", "CPA, ₽", "CR, %"],
                     rows, formats={"Показы": INTEGER, "Клики": INTEGER, "CTR, %": PERCENT,
                                    "Расход, ₽": MONEY, "Конверсии": INTEGER, "CPA, ₽": MONEY, "CR, %": PERCENT},
                     widths=[45, 12, 10, 10, 14, 12, 12, 10])

    # 7. Устройства и регионы
    ws = _sheet(wb, "Устройства и регионы")
    rows = []
    for name, label in (("devices", "Устройства"), ("regions", "Регионы")):
        for r in audit.get("segment_rows", {}).get(name) or []:
            rows.append({"Срез": label, "Сегмент": r.get("key"), "Показы": r.get("impressions"),
                         "Клики": r.get("clicks"), "Расход, ₽": r.get("cost"),
                         "Конверсии": r.get("conversions"), "CPA, ₽": r.get("cpa"), "CR, %": r.get("cr")})
    if not rows:
        _empty(ws, "stats_devices и stats_regions не получены или строк нет")
    else:
        _write_table(ws, ["Срез", "Сегмент", "Показы", "Клики", "Расход, ₽", "Конверсии", "CPA, ₽", "CR, %"],
                     rows, formats={"Показы": INTEGER, "Клики": INTEGER, "Расход, ₽": MONEY,
                                    "Конверсии": INTEGER, "CPA, ₽": MONEY, "CR, %": PERCENT},
                     widths=[14, 35, 12, 10, 14, 12, 12, 10])

    # 8. Рекомендации
    ws = _sheet(wb, "Рекомендации")
    recommendations = audit.get("recommendations") or []
    if not recommendations:
        _empty(ws, "рекомендаций нет при текущих данных")
    else:
        rows = []
        for rec in recommendations:
            effect = rec.get("expected_effect") or {}
            verify = rec.get("verify") or {}
            rows.append({
                "ID": rec.get("id"), "Приоритет": rec.get("priority"),
                "Уверенность": rec.get("confidence"), "Правило": rec.get("rule_id"),
                "Проблема": rec.get("problem"), "Объект": rec.get("object", {}).get("type"),
                "Объект ID": fmt_id(rec.get("object", {}).get("id")),
                "Что сделать": rec.get("proposed_change"),
                "Эффект": effect.get("type"), "Оценка эффекта": effect.get("value"),
                "Метод оценки": effect.get("method"), "Риск": rec.get("risk", {}).get("level"),
                "Способ выполнения": ("инструмент " + str(rec.get("apply", {}).get("tool"))
                                      if rec.get("apply", {}).get("method") == "tool"
                                      else rec.get("apply", {}).get("manual")),
                "Проверка": verify.get("expected_direction"),
                "Успех": verify.get("success"), "Неуспех": verify.get("failure"),
                "Доказательства": "; ".join(
                    f"{e.get('source')}:{e.get('field')}={e.get('value')}" for e in rec.get("evidence") or []),
            })
        _write_table(ws, ["ID", "Приоритет", "Уверенность", "Правило", "Проблема", "Объект", "Объект ID",
                          "Что сделать", "Эффект", "Оценка эффекта", "Метод оценки", "Риск",
                          "Способ выполнения", "Проверка", "Успех", "Неуспех", "Доказательства"],
                     rows, formats={"Оценка эффекта": NUMBER},
                     widths=[8, 10, 12, 10, 40, 12, 22, 55, 12, 14, 35, 8, 30, 30, 30, 30, 55])
        from openpyxl.styles import PatternFill
        for r_index, rec in enumerate(recommendations, start=2):
            color = PRIORITY_FILL.get(rec.get("priority"))
            if color:
                ws.cell(row=r_index, column=2).fill = PatternFill("solid", fgColor=color)

    # 9. Источники и ограничения
    ws = _sheet(wb, "Источники и ограничения")
    _header(ws, ["Тип", "Название", "Детали", "Значение"])
    row_index = 2
    for source in audit.get("data_sources") or []:
        data = [
            ("Источник", source.get("action"), source.get("file"), source.get("rows")),
            ("SHA-256", source.get("action"), "контрольная сумма", source.get("sha256")),
            ("Статус", source.get("action"), source.get("note") or "", source.get("status")),
        ]
        for kind, name, details, value in data:
            ws.cell(row=row_index, column=1, value=kind)
            ws.cell(row=row_index, column=2, value=name)
            ws.cell(row=row_index, column=3, value=details)
            cell = ws.cell(row=row_index, column=4, value=value)
            cell.number_format = TEXT
            row_index += 1
    for limitation in audit.get("limitations") or []:
        ws.cell(row=row_index, column=1, value="Ограничение")
        ws.cell(row=row_index, column=2, value=limitation.get("id"))
        ws.cell(row=row_index, column=3, value=limitation.get("text"))
        ws.cell(row=row_index, column=4, value=str(limitation.get("detail") or ""))
        row_index += 1
    for check in (audit.get("data_quality") or {}).get("checks") or []:
        ws.cell(row=row_index, column=1, value="Проверка данных")
        ws.cell(row=row_index, column=2, value=check.get("id"))
        ws.cell(row=row_index, column=3, value=check.get("title"))
        ws.cell(row=row_index, column=4, value=check.get("status"))
        row_index += 1
    ws.auto_filter.ref = f"A1:D{row_index - 1}"
    _widths(ws, [18, 22, 70, 45])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
